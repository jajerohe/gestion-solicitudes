import unicodedata
import pandas as pd

HOJAS = ["DETALLE_GENERAL", "DETALLE_FUNCIONALES"]

COLUMNAS_OBLIGATORIAS = {
    "ID_SOLICITUD", "PRODUCT_OWNER", "STATUS", "FECHA_APERTURA"
}

COLUMNAS_SQL = [
    "ORIGEN", "TIPO", "ID_SOLICITUD", "TI_PRESTADOR",
    "GRUPO_ASIGNACION", "ASIGNADO_A", "NOMBRE_ASIGNATARIO",
    "CORREO_ASIGNATARIO", "CATEGORIA", "CURRENT_PHASE",
    "CONTACTO", "DESTINATARIO", "STATUS", "FECHA_APERTURA",
    "RANGTIEMPO", "TITULO", "DESCRIPCION", "SUBSERVICIO_AFECTADO",
    "PRODUCT_OWNER", "DIFDIAS"
]

def quitar_acentos(texto):
    texto = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in texto if not unicodedata.combining(c))

def normalizar_nombre_columna(nombre):
    nombre = quitar_acentos(str(nombre).strip()).upper()
    for c in (" ", "/", "-"):
        nombre = nombre.replace(c, "_")
    while "__" in nombre:
        nombre = nombre.replace("__", "_")
    return nombre

def normalizar_columnas(df):
    df = df.copy()
    df.columns = [normalizar_nombre_columna(c) for c in df.columns]
    return df.rename(columns={
        "SUMA_DE_DIFDIAS": "DIFDIAS",
        "SUMA_DE_DIAS": "DIFDIAS",
    })

def validar_columnas(df):
    faltantes = COLUMNAS_OBLIGATORIAS - set(df.columns)
    if faltantes:
        raise ValueError(
            "Columnas obligatorias faltantes: " + ", ".join(sorted(faltantes))
        )

def limpiar_valor(valor):
    if pd.isna(valor):
        return None
    if isinstance(valor, pd.Timestamp):
        return valor.to_pydatetime()
    if hasattr(valor, "to_pydatetime"):
        return valor.to_pydatetime()
    if isinstance(valor, str):
        valor = valor.strip()
        return valor if valor else None
    return valor

def normalizar_product_owner(valor):
    return " ".join(str(valor or "").split()).upper()

def procesar_hoja(df, hoja, product_owners):
    """Filtra la hoja dejando solo las solicitudes de los Product Owner
    recibidos (los NOMBRE de la tabla PODS)."""
    df = normalizar_columnas(df)
    validar_columnas(df)

    permitidos = {normalizar_product_owner(po) for po in product_owners}
    df["PRODUCT_OWNER"] = df["PRODUCT_OWNER"].map(normalizar_product_owner)
    df = df[df["PRODUCT_OWNER"].isin(permitidos)].copy()

    if "FECHA_APERTURA" in df.columns:
        df["FECHA_APERTURA"] = pd.to_datetime(df["FECHA_APERTURA"], errors="coerce")
    if "FECHA_CIERRE" in df.columns:
        df["FECHA_CIERRE"] = pd.to_datetime(df["FECHA_CIERRE"], errors="coerce")
    if "DIFDIAS" in df.columns:
        df["DIFDIAS"] = pd.to_numeric(df["DIFDIAS"], errors="coerce")

    return df

def preparar_registro(fila, hoja):
    valores = []
    for columna in COLUMNAS_SQL:
        if columna == "ORIGEN":
            valores.append(hoja)
        elif columna == "CURRENT_PHASE":
            valores.append(None)
        else:
            valor = limpiar_valor(fila.get(columna))
            # Los títulos se guardan siempre en MAYÚSCULAS.
            if columna == "TITULO" and isinstance(valor, str):
                valor = valor.upper()
            valores.append(valor)
    return tuple(valores)

def leer_hojas(archivo, product_owners):
    """Lee el Excel una sola vez y devuelve [(hoja, df_filtrado), ...] con las
    solicitudes de los Product Owner recibidos (los NOMBRE de la tabla PODS)."""
    excel = pd.ExcelFile(archivo)
    return [
        (hoja, procesar_hoja(pd.read_excel(excel, sheet_name=hoja), hoja, product_owners))
        for hoja in HOJAS
        if hoja in excel.sheet_names
    ]

COLUMNAS_VISTA_PREVIA = [
    "ID_SOLICITUD", "PRODUCT_OWNER", "STATUS", "FECHA_APERTURA", "TITULO",
    "SUBSERVICIO_AFECTADO", "RANGTIEMPO", "DIFDIAS",
]

def analizar_excel(hojas):
    """Vista previa de las solicitudes a cargar a partir de leer_hojas()."""
    partes = []
    for hoja, df in hojas:
        parte = pd.DataFrame({"HOJA": hoja}, index=df.index)
        for columna in COLUMNAS_VISTA_PREVIA:
            parte[columna] = df[columna].map(limpiar_valor) if columna in df.columns else None
        partes.append(parte)
    if not partes:
        return pd.DataFrame(columns=["HOJA", *COLUMNAS_VISTA_PREVIA])
    return pd.concat(partes, ignore_index=True)

SQL_CARGA = '''
    INSERT INTO public."SOLICITUDES"
    ("ORIGEN","TIPO","ID_SOLICITUD","TI_PRESTADOR","GRUPO_ASIGNACION",
     "ASIGNADO_A","NOMBRE_ASIGNATARIO","CORREO_ASIGNATARIO","CATEGORIA",
     "CURRENT_PHASE","CONTACTO","DESTINATARIO","STATUS","FECHA_APERTURA",
     "RANGTIEMPO","TITULO","DESCRIPCION","SUBSERVICIO_AFECTADO",
     "PRODUCT_OWNER","DIFDIAS")
    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,
            %s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
    ON CONFLICT ("ID_SOLICITUD") DO UPDATE
    SET
        "ORIGEN" = EXCLUDED."ORIGEN",
        "TIPO" = EXCLUDED."TIPO",
        "TI_PRESTADOR" = EXCLUDED."TI_PRESTADOR",
        "GRUPO_ASIGNACION" = EXCLUDED."GRUPO_ASIGNACION",
        "ASIGNADO_A" = EXCLUDED."ASIGNADO_A",
        "NOMBRE_ASIGNATARIO" = EXCLUDED."NOMBRE_ASIGNATARIO",
        "CORREO_ASIGNATARIO" = EXCLUDED."CORREO_ASIGNATARIO",
        "CATEGORIA" = EXCLUDED."CATEGORIA",
        "CONTACTO" = EXCLUDED."CONTACTO",
        "DESTINATARIO" = EXCLUDED."DESTINATARIO",
        "STATUS" = EXCLUDED."STATUS",
        "FECHA_APERTURA" = EXCLUDED."FECHA_APERTURA",
        "RANGTIEMPO" = EXCLUDED."RANGTIEMPO",
        "TITULO" = EXCLUDED."TITULO",
        "DESCRIPCION" = EXCLUDED."DESCRIPCION",
        "SUBSERVICIO_AFECTADO" = EXCLUDED."SUBSERVICIO_AFECTADO",
        "PRODUCT_OWNER" = EXCLUDED."PRODUCT_OWNER",
        "DIFDIAS" = EXCLUDED."DIFDIAS"
    WHERE public."SOLICITUDES"."FECHA_CIERRE" IS NULL
      AND public."SOLICITUDES"."CODIGO_CIERRE" IS NULL
    RETURNING "ID_SOLICITUD"
'''

def cargar_excel(hojas, db):
    """Carga las solicitudes de leer_hojas() en public."SOLICITUDES".

    Si el ID ya existe, el UPDATE solo se ejecuta cuando la solicitud sigue
    abierta (FECHA_CIERRE y CODIGO_CIERRE son NULL); las cerradas quedan
    protegidas. Todas las filas de una hoja se envían en un solo lote.
    """
    resultado = {
        "generales": 0, "funcionales": 0, "duplicados": 0,
        "errores": 0, "hojas": []
    }

    for hoja, df in hojas:
        resultado["hojas"].append({"hoja": hoja, "registros": len(df)})
        if df.empty:
            continue

        registros = [preparar_registro(fila, hoja) for _, fila in df.iterrows()]
        db.cursor.executemany(SQL_CARGA, registros, returning=True)

        for _ in db.cursor.results():
            # INSERT nuevo o UPDATE de una solicitud abierta.
            if db.cursor.fetchone():
                if hoja == "DETALLE_GENERAL":
                    resultado["generales"] += 1
                else:
                    resultado["funcionales"] += 1
            else:
                # Sin fila en RETURNING: el ID existía y estaba cerrado.
                resultado["duplicados"] += 1

    return resultado

# ============================================================
# TABLA JDU (la misma del correo de backlog y de la hoja JDU)
# ============================================================
ORDEN_RANGOS = ["0-5 días", "6-10 días", "11-20 días", ">20 días"]


def _orden_rango(rango):
    return ORDEN_RANGOS.index(rango) if rango in ORDEN_RANGOS else len(ORDEN_RANGOS)


def calcular_tabla_jdu(registros):
    """Cuenta las solicitudes por Product Owner, TIPO y RANGTIEMPO.

    registros: [(product_owner, tipo, rango), ...].
    Devuelve (conteo, tipos, filas, total_po):
      conteo[(po, tipo, rango)] -> cantidad
      tipos  {tipo: [rangos en orden 0-5, 6-10, 11-20, >20]}
      filas  Product Owner en el orden del reporte: primero los que tienen
             solicitudes de más de 5 días (de mayor a menor por esa cantidad y
             luego por el total); después los demás, de mayor a menor por el
             total. Empates: orden alfabético.
      total_po {po: total}
    """
    conteo, tipos = {}, {}
    for po, tipo, rango in registros:
        conteo[(po, tipo, rango)] = conteo.get((po, tipo, rango), 0) + 1
        tipos.setdefault(tipo, set()).add(rango)

    tipos = {t: sorted(r, key=_orden_rango) for t, r in sorted(tipos.items())}
    total_po, mas_5_dias = {}, {}
    for (po, _, rango), c in conteo.items():
        total_po[po] = total_po.get(po, 0) + c
        if rango in ORDEN_RANGOS[1:]:      # 6-10, 11-20 y >20 días
            mas_5_dias[po] = mas_5_dias.get(po, 0) + c

    filas = sorted(
        total_po,
        key=lambda po: (
            0 if mas_5_dias.get(po) else 1,
            -mas_5_dias.get(po, 0),
            -total_po[po],
            po,
        )
    )
    return conteo, tipos, filas, total_po


def agregar_hoja_jdu(libro, general, funcional):
    """Agrega (o reemplaza) la hoja JDU, como primera hoja del libro, con las
    mismas tablas del correo de backlog: Detalle General y Grupo Funcional
    Ecopetrol, con el mismo orden y los colores del sistema."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    if "JDU" in libro.sheetnames:
        del libro["JDU"]
    ws = libro.create_sheet("JDU", 0)
    ws.sheet_view.showGridLines = False

    def relleno(color):
        return PatternFill("solid", fgColor=color)

    VERDE_OSCURO, LIMA, FONDO = "004236", "CCD32A", "F4F5F7"
    borde = Border(bottom=Side(style="thin", color="E0E5E5"))
    f_titulo = Font(name="Calibri", size=13, bold=True, color=VERDE_OSCURO)
    f_enc = Font(name="Calibri", size=9, bold=True, color="5F6D69")
    f_tipo = Font(name="Calibri", size=9, bold=True, color=VERDE_OSCURO)
    f_dato = Font(name="Calibri", size=11, color="17342F")
    f_negrita = Font(name="Calibri", size=11, bold=True, color="17342F")
    f_total_tipo = Font(name="Calibri", size=11, bold=True, color=VERDE_OSCURO)
    f_blanca = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    derecha = Alignment(horizontal="right", vertical="bottom", wrap_text=True)
    centro = Alignment(horizontal="center", vertical="center")

    fila = 1
    anchos = {1: 36}

    def escribir_tabla(titulo, registros, fila):
        ws.cell(fila, 1, titulo).font = f_titulo
        fila += 1
        if not registros:
            ws.cell(fila, 1, "Sin solicitudes").font = f_dato
            return fila + 1

        conteo, tipos, filas_po, total_po = calcular_tabla_jdu(registros)
        f1, f2 = fila, fila + 1

        def enc(r, c, valor, fuente=f_enc, fondo=FONDO, alineacion=derecha):
            celda = ws.cell(r, c, valor)
            celda.font, celda.fill, celda.alignment = fuente, relleno(fondo), alineacion
            return celda

        enc(f1, 1, None); enc(f2, 1, "PRODUCT OWNER", alineacion=Alignment(vertical="bottom"))
        col = 2
        columnas = []           # (tipo, rango | None para total del tipo)
        for tipo, rangos in tipos.items():
            ws.merge_cells(start_row=f1, start_column=col, end_row=f1, end_column=col + len(rangos) - 1)
            celda = enc(f1, col, str(tipo).upper(), f_tipo, alineacion=centro)
            for c in range(col, col + len(rangos)):
                ws.cell(f1, c).fill = relleno(FONDO)
                ws.cell(f1, c).border = Border(bottom=Side(style="medium", color=LIMA))
            for r in rangos:
                enc(f2, col, str(r).upper())
                columnas.append((tipo, r)); col += 1
            enc(f1, col, None, fondo="EEF4E6")
            enc(f2, col, f"TOTAL {str(tipo).upper()}", f_tipo, "EEF4E6")
            columnas.append((tipo, None)); col += 1
        enc(f1, col, None); enc(f2, col, "TOTAL GENERAL", f_tipo)
        ultima = col

        fila = f2 + 1
        for po in filas_po:
            ws.cell(fila, 1, po).font = f_dato
            for i, (tipo, r) in enumerate(columnas, start=2):
                celda = ws.cell(fila, i)
                if r is None:
                    valor = sum(conteo.get((po, tipo, x), 0) for x in tipos[tipo])
                    celda.value = valor or None
                    celda.font, celda.fill = f_total_tipo, relleno("EEF4E6")
                else:
                    valor = conteo.get((po, tipo, r), 0)
                    if valor:
                        celda.value = valor
                        if r == ORDEN_RANGOS[0]:
                            celda.font, celda.fill = f_negrita, relleno("F7DB17")
                        else:
                            celda.font, celda.fill = f_blanca, relleno("FF5F00")
            ws.cell(fila, ultima, total_po[po]).font = f_negrita
            for c in range(1, ultima + 1):
                ws.cell(fila, c).border = borde
            fila += 1

        ws.cell(fila, 1, "Total general")
        for i, (tipo, r) in enumerate(columnas, start=2):
            rangos = tipos[tipo] if r is None else [r]
            ws.cell(fila, i, sum(conteo.get((po, tipo, x), 0) for po in filas_po for x in rangos))
        ws.cell(fila, ultima, sum(total_po.values()))
        for c in range(1, ultima + 1):
            ws.cell(fila, c).font, ws.cell(fila, c).fill = f_blanca, relleno(VERDE_OSCURO)

        for c in range(2, ultima + 1):
            anchos[c] = max(anchos.get(c, 0), 13)
        anchos[1] = max(anchos[1], max(len(str(p)) for p in filas_po) + 4)
        return fila + 1

    fila = escribir_tabla("Detalle General", general, fila)
    if funcional:
        fila = escribir_tabla("Grupo Funcional Ecopetrol", funcional, fila + 2)

    for c, ancho in anchos.items():
        ws.column_dimensions[get_column_letter(c)].width = ancho
    libro.active = 0
    for hoja in libro.worksheets:
        hoja.sheet_view.tabSelected = hoja.title == "JDU"


def filtrar_adjunto(contenido, product_owners, tablas_jdu=None):
    """Devuelve una copia del Excel con el filtro de Excel aplicado en
    PRODUCT_OWNER de DETALLE_GENERAL y DETALLE_FUNCIONALES, dejando visibles
    solo los Product Owner recibidos (los NOMBRE de la tabla PODS).

    Las demás hojas, el formato y los datos se conservan: las filas de otros
    Product Owner solo quedan ocultas por el filtro, como al filtrar a mano.
    Con tablas_jdu=(general, funcional) agrega además la hoja JDU.
    """
    import io
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    permitidos = {normalizar_product_owner(po) for po in product_owners}
    libro = load_workbook(io.BytesIO(contenido))

    for hoja in HOJAS:
        if hoja not in libro.sheetnames:
            continue
        ws = libro[hoja]

        # Encabezados en la primera fila (igual que la lectura con pandas).
        columna_po = next(
            (c.column for c in ws[1] if c.value is not None
             and normalizar_nombre_columna(c.value) == "PRODUCT_OWNER"),
            None
        )
        if columna_po is None or ws.max_row < 2:
            continue

        ultima_columna = max(ws.max_column, columna_po)
        rango = f"A1:{get_column_letter(ultima_columna)}{ws.max_row}"

        # Si los datos están en una Tabla de Excel, el filtro va en la tabla.
        tabla = next(
            (t for t in ws.tables.values()
             if t.ref.split(":")[0].rstrip("0123456789") == "A"
             and t.ref.split(":")[0][1:] == "1"),
            None
        )
        if tabla is not None:
            rango = tabla.ref
            filtro = tabla.autoFilter
            if filtro is None:
                from openpyxl.worksheet.filters import AutoFilter
                tabla.autoFilter = filtro = AutoFilter(ref=rango)
            filtro.ref = rango
            ws.auto_filter.ref = None
        else:
            ws.auto_filter.ref = rango
            filtro = ws.auto_filter
        filtro.filterColumn = []

        visibles = set()
        for fila in range(2, ws.max_row + 1):
            valor = ws.cell(row=fila, column=columna_po).value
            if valor is not None and normalizar_product_owner(valor) in permitidos:
                visibles.add(str(valor))
                ws.row_dimensions[fila].hidden = False
            else:
                ws.row_dimensions[fila].hidden = True

        # Valores marcados en el filtro (como en la lista del filtro de Excel).
        filtro.add_filter_column(columna_po - 1, sorted(visibles) or ["(sin coincidencias)"])

    # Hoja JDU con las mismas tablas del correo (general, funcional).
    if tablas_jdu is not None:
        agregar_hoja_jdu(libro, *tablas_jdu)

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
